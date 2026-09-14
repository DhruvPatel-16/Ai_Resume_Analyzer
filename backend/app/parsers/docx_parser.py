import io
from typing import Tuple, Dict, Any

class DOCXParseError(Exception):
    pass

def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extract structured text and metadata from DOCX bytes.
    Extracts body paragraphs as well as text from embedded tables.
    """
    if not file_bytes or len(file_bytes) < 10:
        raise DOCXParseError("DOCX file is empty or corrupted.")

    # DOCX files are ZIP archives starting with PK\x03\x04
    if not file_bytes.startswith(b"PK\x03\x04"):
        raise DOCXParseError("File signature does not match a valid DOCX document.")

    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        
        paragraphs = []
        for p in doc.paragraphs:
            clean_text = p.text.strip()
            if clean_text:
                paragraphs.append(clean_text)

        # Extract table text as well
        table_paragraphs = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_paragraphs.append(" | ".join(row_cells))

        full_text = "\n".join(paragraphs)
        if table_paragraphs:
            full_text += "\n\n" + "\n".join(table_paragraphs)

        full_text = full_text.strip()

        if len(full_text) < 20:
            raise DOCXParseError("Document contains insufficient text for analysis.")

        metadata = {
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables),
            "parser_used": "python-docx",
        }

        return full_text, metadata
    except DOCXParseError:
        raise
    except Exception as e:
        raise DOCXParseError(f"Failed to extract text from DOCX file: {str(e)}")
